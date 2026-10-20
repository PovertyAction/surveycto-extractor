"""PII classification and the value gate, on an invented project (synthetic_project).

The canary test runs the pipeline and both agent readers and asserts that no value
held only by a flagged column reaches anything an agent reads; the review workbook
next to the .dta is the one place they go.
"""

import dataclasses
import json
import sys
import warnings
from pathlib import Path

import pandas as pd
import pytest
import synthetic_project
from conftest import SM
from openpyxl import load_workbook

from surveycto_extractor import pii, survey
from surveycto_extractor.pii import (
    MissingSpacyModel,
    PiiConfig,
    build_analyzer,
    warn_unknown_names,
    withhold_reasons,
)

CONFIG = dataclasses.replace(SM, blacklist=("dob", "secret_num"))


class Spy:
    """Record every text sent to the analyzer."""

    def __init__(self, inner):
        self.inner = inner
        self.texts: list[str] = []

    @property
    def nlp_engine(self):
        return self.inner.nlp_engine

    def analyze(self, text, language, **kwargs):
        self.texts.append(text)
        return self.inner.analyze(text=text, language=language, **kwargs)


def _run(tmp_path, monkeypatch, analyzer, config=CONFIG):
    """Build the project, run extract and vardict on it, return (config, spy)."""
    cfg = synthetic_project.build(tmp_path, config)
    monkeypatch.setitem(sys.modules, "config", cfg)
    from surveycto_extractor.cli import extract, vardict
    from surveycto_extractor.extractors import json_extractor

    for module in (extract, vardict, json_extractor):
        monkeypatch.setattr(module, "config", cfg)
    monkeypatch.setattr(vardict, "DATASETS", cfg.DATASETS)
    spy = Spy(analyzer)
    monkeypatch.setattr(pii, "build_analyzer", lambda _config: spy)
    monkeypatch.setattr(sys, "argv", ["extract", "--survey", "synth"])
    extract.main()
    monkeypatch.setattr(
        sys, "argv", ["vardict", "--survey", "synth", "--xlsx", "--validate"]
    )
    vardict.main()
    return cfg, spy


@pytest.fixture
def project(tmp_path, monkeypatch, analyzer):
    return _run(tmp_path, monkeypatch, analyzer)


def _variables(cfg):
    path = cfg.DATASETS["synth"]["output_json"]
    return json.loads(Path(path).read_text(encoding="utf-8"))["variables"]


def test_every_column_gets_a_block(project):
    cfg, _ = project
    variables = _variables(cfg)
    assert set(variables) == set(synthetic_project.ROWS)
    assert all("pii" in entry for entry in variables.values())


def test_rules(project):
    cfg, _ = project
    blocks = {c: e["pii"] for c, e in _variables(cfg).items()}

    def status(col):
        return blocks[col]["status"], blocks[col]["basis"][0]

    assert blocks["village"] == {"status": "cleared_by_form", "basis": ["publishable"]}
    assert status("age") == ("whitelisted", "type:integer")
    assert status("area") == ("whitelisted", "type:decimal")
    assert status("consent") == ("whitelisted", "choices:inline")
    assert status("starttime") == ("whitelisted", "type:start")
    assert status("key") == ("whitelisted", "system")
    assert status("submissiondate") == ("whitelisted", "system")
    assert status("age2") == ("whitelisted", "lineage:age")
    assert status("elapsed") == ("whitelisted", "generated")
    assert status("dob") == ("contains_pii", "reviewed")
    assert status("secret_num") == ("contains_pii", "reviewed")
    assert status("secret_calc") == ("contains_pii", "lineage:secret_num")
    assert status("deviceid") == ("contains_pii", "device:deviceid")
    assert status("username") == ("contains_pii", "device:username")
    assert status("photo") == ("contains_pii", "type:image")
    assert status("gpslatitude") == ("contains_pii", "type:geopoint")
    assert status("gpsaccuracy") == ("contains_pii", "type:geopoint")
    assert status("lat") == ("contains_pii", "lineage:gps")
    assert status("coder") == ("contains_pii", "calls:enumerator-name()")
    assert blocks["respondent"]["basis"] == ["type:text", "presidio:PERSON"]
    assert blocks["respondent"]["presidio"]["hit_rows"] == 3
    assert "presidio:PHONE_NUMBER" in blocks["phone"]["basis"]
    assert blocks["tag"]["status"] == "may_contain_pii"
    assert blocks["prev"]["basis"][:2] == ["type:calculate", "pulldata"]
    assert blocks["hhid"]["status"] == "no_hits"
    assert "presidio:PHONE_NUMBER" in blocks["phone_num"]["basis"]
    assert status("enum_user") == ("contains_pii", "metadata:username")
    # Selects whose choices come from data are scanned, not cleared as inline.
    assert blocks["hh_pick"]["status"] == "no_hits"
    assert blocks["roster_pick"]["status"] == "no_hits"


def test_publishable_is_never_scanned(project):
    _, spy = project
    assert "Villa Ficticia Uno" not in spy.texts


def test_blacklisted_is_never_scanned(project):
    _, spy = project
    assert "987654" not in spy.texts


def test_only_shown_columns_keep_a_range(project):
    cfg, _ = project
    variables = _variables(cfg)
    assert variables["age"]["data_min"] == 29
    assert variables["area"]["data_max"] == 3.0
    for col in ("secret_num", "secret_calc", "dob", "phone_num"):
        assert "data_min" not in variables[col]
        assert "data_max" not in variables[col]


def test_external_list_is_withheld(project):
    cfg, _ = project
    out = cfg.SURVEYS["synth"]["output_dir"]
    questions = json.loads((out / "synth_questions.json").read_text(encoding="utf-8"))
    hh_pick = next(q for q in questions if q["variable_name"] == "hh_pick")
    assert hh_pick["choices"] is None
    assert hh_pick["choices_withheld"] == ["presidio:PERSON"]


def test_search_list_read_by_the_xml_overlay_is_withheld(project):
    cfg, _ = project
    survey = _variables(cfg)["roster_pick"]["survey"]
    assert survey["choices_withheld"] == ["presidio:PERSON"]
    assert survey["choices"] == [{"value": "value", "label": "label"}]


def test_whitelist_by_stata_name_shows_the_range(tmp_path, monkeypatch, analyzer):
    config = dataclasses.replace(CONFIG, whitelist=("phone_num", "gpsaccuracy"))
    cfg, _ = _run(tmp_path, monkeypatch, analyzer, config)
    variables = _variables(cfg)
    assert variables["phone_num"]["pii"] == {
        "status": "whitelisted",
        "basis": ["reviewed"],
    }
    assert variables["phone_num"]["data_min"] == 2125550147
    assert variables["gpsaccuracy"]["pii"]["status"] == "whitelisted"


def _leaked(texts, canaries=synthetic_project.CANARIES):
    seen = "\n".join(t for t in texts if t)
    return [c for c in canaries if c in seen]


# Tool and shell answers: every canary, and the withheld codes in column names.
ANSWER_CANARIES = synthetic_project.CANARIES + synthetic_project.CODE_CANARIES


def test_no_flagged_value_reaches_an_agent(capsys, project):
    """The output folder, search_survey.py's answers and the console (capsys is set
    up before the project, so it holds the pipeline's own output too).
    """
    cfg, _ = project
    texts = []
    for path in sorted(cfg.SURVEYS["synth"]["output_dir"].rglob("*")):
        if path.suffix in (".json", ".csv", ".txt"):
            texts.append(path.read_text(encoding="utf-8"))
        elif path.suffix == ".xlsx":
            for sheet in load_workbook(path, read_only=True).worksheets:
                for row in sheet.iter_rows(values_only=True):
                    texts.append(" ".join(str(v) for v in row if v is not None))
    pipeline = capsys.readouterr()
    assert "PII: " in pipeline.out  # the pipeline's own output was captured
    assert _leaked([*texts, pipeline.out, pipeline.err]) == []

    # The skill's search_survey.py is survey.py; its shell front end, every command.
    shell = ["--dir", str(cfg.SURVEYS["synth"]["output_dir"])]
    names = list(synthetic_project.ROWS)
    survey.main([*shell, "info"])  # discovers the folder itself
    store = survey.SurveyStore(survey.discover(cfg.SURVEYS["synth"]["output_dir"]))
    for name in names:
        survey.main([*shell, "var", name, "--context", "3"], store=store)
        survey.main([*shell, "gates", name], store=store)
        survey.main([*shell, "neighborhood", name, "--depth", "2"], store=store)
    survey.main([*shell, "var", *names], store=store)
    for words in ("name", "phone", "household roster"):
        survey.main([*shell, "search", words], store=store)
    for list_name in ("yesno", "hh", "roster", "staff"):
        survey.main([*shell, "choices", list_name], store=store)
    survey.main([*shell, "search", "staff", "visit"], store=store)
    survey.main([*shell, "repeats"], store=store)
    survey.main([*shell, "info"], store=store)

    answers = capsys.readouterr()
    assert "staff_multi_<1>" in answers.out
    assert _leaked([answers.out, answers.err], ANSWER_CANARIES) == []
    # The review workbook next to the .dta is where they go.
    review = Path(cfg.DATASETS["synth"]["data"]).with_name("synth_pii_review.xlsx")
    values = {
        str(row[2])
        for row in load_workbook(review, read_only=True)["Values"].iter_rows(
            min_row=2, values_only=True
        )
    }
    assert {"Robert Brown", "dev-2"} <= values


def test_no_flagged_value_in_mcp_answers(project, monkeypatch):
    from surveycto_extractor.mcp_server import survey_server as mcp

    cfg, _ = project
    monkeypatch.setattr(mcp, "_store", survey.SurveyStore(cfg.DATASETS))
    names = list(synthetic_project.ROWS)
    texts = []
    for name in names:
        texts.append(mcp.lookup_variable(name, context=3))
        texts.append(mcp.get_gate_chain(name))
        texts.append(mcp.get_variable_neighborhood(name, depth=2))
    texts.append(mcp.lookup_variables(names))
    texts.append(mcp.get_survey_info())
    texts.append(mcp.get_repeat_structure())
    for words in ("name", "phone", "household roster", "date of birth", "staff"):
        texts.append(mcp.search_questions(words))
    for list_name in ("yesno", "hh", "roster", "staff"):
        texts.append(mcp.get_choice_list(list_name))
    assert _leaked(texts, ANSWER_CANARIES) == []


def test_search_list_withheld_with_its_decoded_labels(analyzer):
    from surveycto_extractor.cli.vardict import _screen_search_lists

    source = {"kind": "search", "dataset": "roster", "mode": "matches", "filter": {}}
    variables = {
        f"pick_{code}": {
            "survey": {
                "original_variable_name": "pick",
                "choice_list": "roster",
                "choices": [
                    {"value": "1", "label": "John Smith"},
                    {"value": "2", "label": "Mary Johnson"},
                ],
            },
            "contract": {
                "data_source": source,
                "choice_code": code,
                "choice_label": label,
            },
        }
        for code, label in (("1", "John Smith"), ("2", "Mary Johnson"))
    }
    literal = [{"value": "hhid", "label": "name"}]
    questions = [{"variable_name": "pick", "choices": literal}]
    _screen_search_lists(variables, questions, SM, analyzer)
    for entry in variables.values():
        assert entry["survey"]["choices"] == literal
        assert entry["survey"]["choices_withheld"] == ["presidio:PERSON"]
        assert "choice_label" not in entry["contract"]
    assert "John Smith" not in json.dumps(variables)


def test_withhold_reasons_order(analyzer):
    many = [{"value": str(i), "label": f"c{i}"} for i in range(51)]
    assert withhold_reasons({"q"}, many, SM, None) == ["51 choices"]
    config = dataclasses.replace(SM, blacklist=("q",))
    assert withhold_reasons({"q"}, many, config, None) == ["reviewed"]
    # The team's review outranks the size cap.
    config = dataclasses.replace(SM, whitelist=("q",))
    assert withhold_reasons({"q"}, many, config, None) == []
    # A list no question uses is still scanned.
    names = [{"value": "1", "label": "John Smith"}]
    assert withhold_reasons(set(), names, config, analyzer) == ["presidio:PERSON"]


def _rules(questions, config):
    q = {"calculation": None, "choices_from_data": False, "publishable": False}
    return pii._Rules(
        [{**q, **x} for x in questions],
        set(config.whitelist),
        set(config.blacklist),
        config.trust_publishable,
    )


def test_blacklisted_dataset_column_carries_through_pulldata():
    questions = [
        {
            "variable_name": "prev",
            "type": "calculate",
            "calculation": "if(${a}, pulldata('r', 'x', 'k', ${a}), "
            "pulldata('prev_round', 'owner', 'hhid', ${a}))",
        }
    ]
    config = PiiConfig(blacklist=("prev_round.owner",))
    assert _rules(questions, config).carried("prev") == "prev_round.owner"
    assert _rules(questions, PiiConfig()).carried("prev") is None


def test_blacklist_outranks_publishable():
    questions = [{"variable_name": "v", "type": "text", "publishable": True}]
    assert _rules(questions, PiiConfig()).static("v")["status"] == "cleared_by_form"
    config = PiiConfig(blacklist=("v",))
    assert _rules(questions, config).static("v") == {
        "status": "contains_pii",
        "basis": ["reviewed"],
    }


def test_non_numeric_strings_are_counted_not_printed(tmp_path, capsys):
    from surveycto_extractor.generators.load_survey_metadata import (
        filter_truly_numeric,
    )

    pq = tmp_path / "d.parquet"
    pd.DataFrame({"score": ["1", "2", "Robert Brown"]}).to_parquet(pq)
    universe = pd.DataFrame({"stata_type": ["str12"]}, index=["score"])
    _, bad = filter_truly_numeric(universe, str(tmp_path / "d.dta"), str(pq))
    assert bad == ["score"]
    assert "Robert Brown" not in capsys.readouterr().out


def test_config_table():
    config = PiiConfig.from_table({"whitelist": ["a"], "languages": {"en": "x"}})
    assert config.whitelist == ("a",)
    with pytest.raises(ValueError, match="unknown key"):
        PiiConfig.from_table({"whitelst": ["a"]})
    with pytest.raises(ValueError, match="both"):
        PiiConfig(whitelist=("a",), blacklist=("a",))


def test_missing_model_names_the_install_command():
    with pytest.raises(MissingSpacyModel, match="uv add https://.*xx_no_such_model"):
        build_analyzer(PiiConfig(languages={"xx": "xx_no_such_model"}))


def test_unknown_name_warns():
    config = PiiConfig(whitelist=("ghost", "age"), blacklist=("roster.ghost",))
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        warn_unknown_names(config, {"age"})
    assert [str(w.message) for w in caught] == [
        "`ghost` cannot be found",
        "`roster.ghost` cannot be found",
    ]


def test_no_personal_data_keeps_the_blacklist_and_metadata(tmp_path, monkeypatch):
    config = dataclasses.replace(CONFIG, no_personal_data=True)
    cfg, _ = _run(tmp_path, monkeypatch, None, config)
    blocks = {c: e["pii"] for c, e in _variables(cfg).items()}
    assert blocks["respondent"] == {"status": "whitelisted", "basis": ["project"]}
    assert blocks["secret_num"]["basis"] == ["reviewed"]
    assert blocks["deviceid"]["status"] == "contains_pii"
    assert blocks["coder"]["status"] == "contains_pii"
    # Geo is whitelisted in a project without personal data.
    assert blocks["lat"] == {"status": "whitelisted", "basis": ["project"]}


def test_a_crash_in_the_gate_leaves_nothing_unscreened(tmp_path, monkeypatch, analyzer):
    def crash(*args, **kwargs):
        raise RuntimeError("classify failed")

    monkeypatch.setattr(pii, "classify", crash)
    with pytest.raises(RuntimeError, match="classify failed"):
        _run(tmp_path, monkeypatch, analyzer)
    written = (tmp_path / "out" / "synth_variable_dictionary.json").read_text(
        encoding="utf-8"
    )
    # enrich read the roster CSV, but its dictionary was never written.
    assert "Emily Davis" not in written
    assert "data_min" not in written


def test_free_text_system_columns_are_scanned(project):
    cfg, _ = project
    block = _variables(cfg)["instancename"]["pii"]
    assert block["basis"] == ["no_question", "presidio:PERSON"]


def test_pulldata_found_whatever_its_spelling():
    questions = [
        {"variable_name": "a", "type": "integer"},
        {
            "variable_name": "prev",
            "type": "calculate",
            "calculation": "PULLDATA ('roster', 'name', 'hhid', ${a})",
        },
    ]
    rules = _rules(questions, PiiConfig(blacklist=("roster.name",)))
    assert rules.static("prev") is None
    assert rules.carried("prev") == "roster.name"


def test_review_keeps_odd_text_as_text(project):
    cfg, _ = project
    review = Path(cfg.DATASETS["synth"]["data"]).with_name("synth_pii_review.xlsx")
    values = [
        row[2]
        for row in load_workbook(review)["Values"].iter_rows(
            min_row=2, values_only=True
        )
    ]
    assert "=Lucy Gray" in values
